import os
import requests
from openpyxl import Workbook, load_workbook
from datetime import datetime
import pytz
import re

# Настройки
per_page = 200
photo_count = 0

def read_progress():
    if not os.path.exists("progress.txt"):
        with open("progress.txt", "w", encoding="utf-8") as f:
            f.write("page=1\n")
        return 1

    with open("progress.txt", "r", encoding="utf-8") as f:
        lines = f.readlines()
    progress = {}
    for line in lines:
        if "=" in line:
            key, value = line.strip().split("=", 1)
            progress[key] = int(value)
    return progress.get("page")

def save_progress(page):
    with open("progress.txt", "w", encoding="utf-8") as f:
        f.write(f"page={page}\n")

def write_status_to_excel(path, status, row):
    wb = load_workbook(path)
    ws = wb.active

    ws.cell(row=row, column=4).value = status

    wb.save(path)
    wb.close()

def find_first_unprocessed_objects(path):
    wb = load_workbook(path)
    ws = wb.active

    for row in range(1, ws.max_row + 1):
        cell_d = ws[f'D{row + 1}']
        if cell_d.value is None:
            cell_a = ws[f'A{row + 1}']
            wb.close()
            return cell_a.value
        else:
            print(ws[f'C{row + 1}'].value, "обработанно")
    wb.close()
    return None

def get_taxon_id(number_cell, column, excel_objects_path):
    wb = load_workbook(excel_objects_path)
    ws = wb.active
    scientific_name = ws[f'{column}{number_cell}'].value
    wb.close()

    if scientific_name == None:
        return "конец"
    url = "https://api.inaturalist.org/v1/taxa"
    params = {
        "q": scientific_name,
        "per_page": 1
    }
    response = requests.get(url, params=params)
    response.raise_for_status()
    results = response.json().get("results", [])
    return results[0].get("id") if results else None

def write_to_excel(excel_file_path, data_for_excel):
    if os.path.exists(excel_file_path):
        wb = load_workbook(excel_file_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        headers = ['Идентификатор ОФФ', 'Название на русском', 'Путь к файлу', 'Автор', 'Права', 'Место съёмки', 'Координаты съёмки',
                   'Время съёмки', 'время загрузки', 'другие сайты с этим фото', 'Проекты', 'Уровень качества', 'id наблюдения']
        ws.append(headers)

    row = [
        data_for_excel.get('Идентификатор ОФФ', ''),
        data_for_excel.get('Название на русском', ''),
        data_for_excel.get('Путь к файлу', ''),
        data_for_excel.get('Автор', ''),
        data_for_excel.get('Права', ''),
        data_for_excel.get('Место съёмки', ''),
        data_for_excel.get('Координаты съёмки', ''),
        data_for_excel.get('Время съёмки', ''),
        data_for_excel.get('время загрузки', ''),
        data_for_excel.get('другие сайты с этим фото', ''),
        data_for_excel.get('Проекты', ''),
        data_for_excel.get('Уровень качества', ''),
        data_for_excel.get('id наблюдения', '')
    ]
    ws.append(row)
    wb.save(excel_file_path)
    wb.close()

def decimal_to_dms(lat, lon):
    def convert(coord, is_lat):
        direction = ''
        if is_lat:
            direction = 'N' if coord >= 0 else 'S'
        else:
            direction = 'E' if coord >= 0 else 'W'

        coord = abs(coord)
        degrees = int(coord)
        minutes_full = (coord - degrees) * 60
        minutes = int(minutes_full)
        seconds = (minutes_full - minutes) * 60

        return f"{degrees}°{minutes:02d}'{seconds:04.1f}\"{direction}"

    lat_str = convert(lat, is_lat=True)
    lon_str = convert(lon, is_lat=False)
    return f"{lat_str}, {lon_str}"

def parse_time_safe(timestr):
    irkutsk_tz = pytz.timezone("Asia/Irkutsk")
    if not timestr:
        return ""
    try:
        dt = datetime.fromisoformat(timestr).astimezone(irkutsk_tz)
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return ""

def get_observation_info(observation_id):
    try:
        url = f'https://api.inaturalist.org/v1/observations/{observation_id}'
        params = {
            'include_new_projects': "true",
            'locale': 'ru',
            'ttl': '-1',
        }

        response = requests.get(url, params=params, timeout=10)
        response.raise_for_status()
        data = response.json()

        obs = data["results"][0]

        russian_name = obs["taxon"].get("preferred_common_name", "")

        time_observed_raw = obs.get("time_observed_at", "")
        time_created_raw = obs.get("created_at", "")

        time_observed = parse_time_safe(time_observed_raw)
        time_created_at = parse_time_safe(time_created_raw)

        place_photo = obs.get("place_guess", "")
        coord_str = obs.get("location", "")

        # Преобразование координат в DMS
        try:
            lat, lon = map(float, coord_str.split(','))
            coord = decimal_to_dms(lat, lon)
        except Exception:
            coord = ""

        # Сбор ссылок
        links_list = []
        outlinks = obs.get("outlinks", [])
        for item in outlinks:
            url = item.get("url")
            if url:
                links_list.append(url)
        links = "\n ".join(links_list)

        # Извлечение проектов
        projects = []

        # Обычные проекты
        for p in obs.get("project_observations", []):
            project = p.get("project", {})
            slug = project.get("slug")
            if slug:
                projects.append(f"https://www.inaturalist.org/projects/{slug}")

        # Нетрадиционные проекты
        for p in obs.get("non_traditional_projects", []):
            project = p.get("project", {})
            slug = project.get("slug")
            if slug:
                projects.append(f"https://www.inaturalist.org/projects/{slug}")

        # Убираем дубли и объединяем
        project_links = "\n ".join(sorted(set(projects)))

        observation_info = {
            'Название на русском': russian_name,
            'Место съёмки': place_photo,
            'Координаты съёмки': coord,
            'Время съёмки': time_observed,
            'время загрузки': time_created_at,
            'другие сайты с этим фото': links,
            'Проекты': project_links,
        }
        return observation_info

    except Exception as e:
        print(f"Ошибка при получении информации о наблюдении {observation_id}: {e}")
        return {
            'Название на русском': '',
            'Место съёмки': '',
            'Координаты съёмки': '',
            'Время съёмки': '',
            'время загрузки': '',
            'другие сайты с этим фото': '',
            'Проекты': '',
        }

def process_photo_data(photo, result, observation_info, main_folder_path, main_name, photo_count):
    try:
        img_url = photo.get('url', '').replace('square', 'medium')
        license_code = photo.get('license_code', 'нет лицензии')
        taxon_name = result.get('taxon', {}).get('preferred_common_name', main_name)
        photo_id = photo.get('id','')
        quality_grade = result.get('quality_grade', '')

        raw_attribution = photo.get('attribution', '')
        cleaned_attribution = re.sub(r'\(c\) ', '', raw_attribution)
        cleaned_attribution = re.split(r', some rights reserved ', cleaned_attribution)[0]

        if 'all rights reserved' in raw_attribution.lower():
            print(f"Пропущено фото № {photo_count} — all rights reserved")
            return

        attribution = cleaned_attribution

        # Папка для сохранения
        folder_path = os.path.join(main_folder_path, main_name, taxon_name)
        if not os.path.exists(folder_path):
            os.makedirs(folder_path)

        file_name = os.path.join(folder_path, f'{photo_id}.jpg')
        if os.path.exists(file_name):
            print(f'Фото № {photo_count} уже загружено')
            return

        # Загрузка изображения
        img_data = requests.get(img_url, timeout=10).content
        with open(file_name, 'wb') as f:
            f.write(img_data)

        print(f'Фото № {photo_count} расположено в папке {folder_path}')

        # Сохраняем в Excel
        data_for_excel = {
            'Идентификатор ОФФ': taxon_name,
            'Название на русском': observation_info.get('Название на русском', ''),
            'Путь к файлу': f'data/images/flora/{main_name}/{taxon_name}/{photo_id}.jpg',
            'Автор': attribution,
            'Права': license_code,
            'Место съёмки': observation_info.get('Место съёмки', ''),
            'Координаты съёмки': observation_info.get('Координаты съёмки', ''),
            'Время съёмки': observation_info.get('Время съёмки', ''),
            'время загрузки': observation_info.get('время загрузки', ''),
            'другие сайты с этим фото': observation_info.get('другие сайты с этим фото', ''),
            'Проекты': observation_info.get('Проекты', ''),
            'Уровень качества': quality_grade,
            'id наблюдения': result.get('id')
        }
        write_to_excel(excel_file_path, data_for_excel)

    except Exception as e:
        print(f"Ошибка при обработке фото № {photo_count}: {e}")


project_dir = os.path.dirname(os.path.abspath(__file__))
main_folder_path = os.path.join(project_dir, 'objects')
if not os.path.exists(main_folder_path):
    os.makedirs(main_folder_path)
excel_file_path = os.path.join(main_folder_path, 'All_Objects.xlsx')
excel_objects_path = os.path.join(project_dir, 'OFF_actual.xlsx')
number_cell = find_first_unprocessed_objects(excel_objects_path)

print("Выберите режим:")
print("1 — Загрузка фото и заполнение Excel")
print("2 — Проверка и заполнение Excel (без загрузки фото)")
mode = input().strip()

try:
    if mode == "1":
        while True:
            number_cell += 1

            taxon_id = get_taxon_id(number_cell, "C", excel_objects_path)
            if taxon_id is None:
                taxon_id = get_taxon_id(number_cell, "B", excel_objects_path)
            if taxon_id == "конец":
                break
            page = read_progress()
            if taxon_id is not None:
                while True:
                    print(f'Запрашиваем страницу {page}...')

                    url = 'https://api.inaturalist.org/v1/observations'
                    params = {
                        'taxon_id': taxon_id,
                        'order_by': 'votes',
                        'quality_grade': 'research',
                        'photos': 'true',
                        'page': page,
                        'per_page': per_page
                    }

                    try:
                        response = requests.get(url, params=params, timeout=10)
                        response.raise_for_status()
                        data = response.json()
                        main_name = data['results'][0]['taxon']['name']
                    except Exception as e:
                        print(f'Ошибка при запросе: {e}')
                        break

                    results = data.get('results', [])
                    if not results:
                        print('Все фото загружены.')
                        break

                    for result in data.get('results', []):
                        observation_id = result.get('id')
                        print(observation_id)
                        observation_info = get_observation_info(observation_id)

                        for photo in result.get('photos', []):
                            photo_count += 1
                            process_photo_data(photo, result, observation_info, main_folder_path, main_name, photo_count)

                    page += 1
                    save_progress(page)
                page = 1
                save_progress(page)
                write_status_to_excel(excel_objects_path, "Обработанно", number_cell)
            else:
                write_status_to_excel(excel_objects_path, "Не найденно", number_cell)
    elif mode == "2":
        wb = load_workbook(excel_file_path)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        for i,row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if None in row:
                row_dict = dict(zip(headers, row))
                observation_info = get_observation_info(str(row[-1]))
                print(row_dict.get('Путь к файлу',''))
                for field in headers:
                    column_index = headers.index(field) + 1
                    current_value = ws.cell(row=i, column=column_index).value
                    if current_value is None:
                        ws.cell(row=i, column=column_index).value = observation_info.get(field, '')
                wb.save(excel_file_path)
        wb.close()
except KeyboardInterrupt:
    print("\nРабота была прервана пользователем.")
except Exception as e:
    print(f"\nПроизошла непредвиденная ошибка: {e}")
finally:
    if mode == "1":
        save_progress(page)
    if mode == "2" and 'wb' in locals():
        wb.close()
        print("Файл Excel закрыт.")